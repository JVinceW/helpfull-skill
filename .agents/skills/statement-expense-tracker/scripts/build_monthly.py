#!/usr/bin/env python3
"""Build the expense workbook for one statement from a normalized statement JSON (schema v1).

Usage:
  python build_monthly.py data/statement_2026-08.json --config expense-config.json --out-dir .
      [--allow-unreconciled] [--uncategorized-to-other] [--force]

Exit codes: 0 ok, 1 invalid input, 2 totals do not reconcile, 3 uncategorized transactions,
4 output already exists (use --force). On success prints a JSON summary to stdout.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import expense_common as ec  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, PieChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402


def build(stmt, cfg, issues, out_path):
    L = cfg["L"]
    money = ec.money_format(cfg)

    def tx(column):
        return ec.tx_range(cfg, column)

    spend, pay, refund = (ec.q(L[key]) for key in ("type_spend", "type_payment", "type_refund"))

    wb = Workbook()
    dash = wb.active
    dash.title = L["dashboard"]
    ws_t = wb.create_sheet(L["transactions"])
    ws_b = wb.create_sheet(L["budget"])
    ws_m = wb.create_sheet(L["merchants"])
    ws_s = wb.create_sheet(L["statement"])

    ec.write_transactions(ws_t, [stmt], cfg)

    # Budget: one row per configured category.
    cats = cfg["categories"]
    ec.header_row(ws_b, 1, [L["col_category"], L["col_budget"], L["col_actual"], L["col_remaining"], L["col_used"]])
    last = len(cats) + 1
    for r, cat in enumerate(cats, start=2):
        ec.put(ws_b, r, 1, cat["name"], literal=True)
        ec.put(ws_b, r, 2, cat.get("budget") or 0, money)
        ec.put(ws_b, r, 3, f"=SUMIFS({tx('net')},{tx('category')},$A{r})", money)
        ec.put(ws_b, r, 4, f"=B{r}-C{r}", money)
        ec.put(ws_b, r, 5, f"=IFERROR(C{r}/B{r},0)", ec.PCT)
    ec.band(ws_b, 2, last, 1, 5)
    total = last + 1
    ec.total_row(ws_b, total,
                 [L["total"], f"=SUM(B2:B{last})", f"=SUM(C2:C{last})", f"=SUM(D2:D{last})", f"=IFERROR(C{total}/B{total},0)"],
                 [None, money, money, money, ec.PCT])
    ec.flag_over(ws_b, f"E2:E{total}", 1)
    ec.flag_negative(ws_b, f"D2:D{total}")
    ec.set_widths(ws_b, {"A": 26, "B": 18, "C": 18, "D": 18, "E": 10})
    ws_b.freeze_panes = "A2"

    # Merchants: groups ordered by net spend.
    groups = ec.sorted_groups([stmt])
    ec.header_row(ws_m, 1, [L["col_group"], L["col_txns"], L["col_net"], L["col_share"]])
    m_last = len(groups) + 1
    for r, group in enumerate(groups, start=2):
        ec.put(ws_m, r, 1, group, literal=True)
        ec.put(ws_m, r, 2, f"=COUNTIFS({tx('group')},$A{r},{tx('type')},{spend})", "#,##0")
        ec.put(ws_m, r, 3, f"=SUMIFS({tx('net')},{tx('group')},$A{r})", money)
        ec.put(ws_m, r, 4, f"=IFERROR(C{r}/SUM($C$2:$C${m_last}),0)", ec.PCT)
    if groups:
        ec.band(ws_m, 2, m_last, 1, 4)
    ec.set_widths(ws_m, {"A": 30, "B": 12, "C": 18, "D": 10})
    ws_m.freeze_panes = "A2"

    # Statement summary.
    status = L["reconciled"] if not issues else L["unreconciled"].replace("{issues}", "; ".join(issues))
    date_format = cfg["date_format"]
    fields = [
        ("f_period", stmt["period"], "@"),
        ("f_institution", stmt["institution"], None),
        ("f_account", stmt["account_label"], None),
        ("f_currency", stmt["currency"] or cfg["currency"].get("code", ""), None),
        ("f_statement_date", stmt["statement_date"], date_format),
        ("f_due_date", stmt["due_date"], date_format),
        ("f_opening", stmt["opening_balance"], money),
        ("f_debits", stmt["total_debits"], money),
        ("f_credits", stmt["total_credits"], money),
        ("f_closing", stmt["closing_balance"], money),
        ("f_min_payment", stmt["minimum_payment"], money),
        ("f_credit_limit", stmt["credit_limit"], money),
        ("f_reconciliation", status, None),
        ("f_generated", dt.datetime.now().replace(microsecond=0), "yyyy-mm-dd hh:mm"),
    ]
    ec.header_row(ws_s, 1, [L["field"], L["value"]])
    rows = {}
    for r, (key, value, fmt) in enumerate(fields, start=2):
        ec.put(ws_s, r, 1, L[key], literal=True)
        ec.put(ws_s, r, 2, value, fmt, literal=True)
        rows[key] = r
    ec.band(ws_s, 2, len(fields) + 1, 1, 2)
    ec.set_widths(ws_s, {"A": 26, "B": 60})

    # Dashboard: KPI cards in A:D, charts below.
    S, B = ec.sheet_ref(L["statement"]), ec.sheet_ref(L["budget"])
    subtitle = " · ".join(x for x in (stmt["institution"], stmt["account_label"], f"{L['col_period']} {stmt['period']}") if x)
    ec.title_block(dash, cfg.get("workbook_title") or L["monthly_title"], subtitle, 4)
    count = f"COUNTIFS({tx('type')},{spend})"

    def stmt_cell(key):
        ref = f"{S}!$B${rows[key]}"
        return f'=IF({ref}="","",{ref})'

    ec.kpi_row(dash, 5, [
        (L["kpi_net_spend"], f"=SUM({tx('net')})", money),
        (L["kpi_credits"], f"=SUMIFS({tx('amount')},{tx('type')},{pay})+SUMIFS({tx('amount')},{tx('type')},{refund})", money),
        (L["kpi_txns"], f"={count}", "#,##0"),
        (L["kpi_avg_txn"], f"=IFERROR(SUMIFS({tx('net')},{tx('type')},{spend})/{count},0)", money),
    ])
    ec.kpi_row(dash, 8, [
        (L["kpi_closing"], stmt_cell("f_closing"), money),
        (L["kpi_min_payment"], stmt_cell("f_min_payment"), money),
        (L["kpi_due"], stmt_cell("f_due_date"), date_format),
        (L["kpi_top_category"], f"=INDEX({B}!$A$2:$A${last},MATCH(MAX({B}!$C$2:$C${last}),{B}!$C$2:$C${last},0))", "@"),
    ])

    pie = PieChart()
    pie.title = L["chart_category"]
    pie.add_data(Reference(ws_b, min_col=3, min_row=1, max_row=last), titles_from_data=True)
    pie.set_categories(Reference(ws_b, min_col=1, min_row=2, max_row=last))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.width, pie.height = 11.5, 8.5
    dash.add_chart(pie, "A11")

    bars = BarChart()
    bars.type = "col"
    bars.title = L["chart_budget"]
    bars.add_data(Reference(ws_b, min_col=2, max_col=3, min_row=1, max_row=last), titles_from_data=True)
    bars.set_categories(Reference(ws_b, min_col=1, min_row=2, max_row=last))
    bars.y_axis.numFmt = "#,##0"
    bars.width, bars.height = 11.5, 8.5
    dash.add_chart(bars, "C11")

    if groups:
        top = min(10, len(groups))
        merchants = BarChart()
        merchants.type = "bar"
        merchants.title = L["chart_merchants"]
        merchants.add_data(Reference(ws_m, min_col=3, min_row=1, max_row=top + 1), titles_from_data=True)
        merchants.set_categories(Reference(ws_m, min_col=1, min_row=2, max_row=top + 1))
        merchants.x_axis.scaling.orientation = "maxMin"
        merchants.legend = None
        merchants.width, merchants.height = 23, 10
        dash.add_chart(merchants, "A30")

    ec.set_widths(dash, {letter: 32 for letter in "ABCD"})
    dash.sheet_view.showGridLines = False
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    ec.utf8_stdio()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("statement", help="normalized statement JSON (schema v1)")
    parser.add_argument("--config", help="expense-config.json (defaults apply when omitted)")
    parser.add_argument("--out-dir", default=".", help="folder for the workbook (default: current folder)")
    parser.add_argument("--allow-unreconciled", action="store_true", help="build even if totals do not match")
    parser.add_argument("--uncategorized-to-other", action="store_true", help="put unmatched transactions in other_category")
    parser.add_argument("--force", action="store_true", help="overwrite an existing workbook")
    args = parser.parse_args()

    try:
        cfg = ec.load_config(args.config)
        stmt = ec.load_statement(args.statement)
        missing = ec.resolve(stmt, cfg, args.uncategorized_to_other)
    except ec.InputError as exc:
        ec.fail(ec.EXIT_BAD_INPUT, f"INVALID INPUT: {exc}")
    if missing:
        lines = [f"  #{t['no']} {t['transaction_date']} {t['description']} {t['amount']}" for t in missing]
        ec.fail(ec.EXIT_UNCATEGORIZED, f"UNCATEGORIZED ({len(missing)}): set 'category' in the JSON or add config rules\n" + "\n".join(lines))
    issues = ec.reconcile(stmt, cfg)
    if issues and not args.allow_unreconciled:
        ec.fail(ec.EXIT_UNRECONCILED, "UNRECONCILED:\n  " + "\n  ".join(issues))

    out_path = Path(args.out_dir) / f"{cfg['file_prefix']}_{stmt['period']}.xlsx"
    if out_path.exists() and not args.force:
        ec.fail(ec.EXIT_EXISTS, f"EXISTS: {out_path} (re-run with --force to overwrite)")
    build(stmt, cfg, issues, out_path)

    result = {"workbook": str(out_path), "reconciliation": issues or "ok", **ec.summary([stmt], cfg)}
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
