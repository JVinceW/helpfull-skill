#!/usr/bin/env python3
"""Build the all-statements expense overview from normalized statement JSON files (schema v1).

Usage:
  python build_master.py data --config expense-config.json --out-dir . [--force]
      [--allow-unreconciled] [--uncategorized-to-other]

Inputs may be JSON files, folders (every *.json inside), or glob patterns. One statement per period.
Exit codes: 0 ok, 1 invalid input, 2 totals do not reconcile, 3 uncategorized transactions,
4 output already exists (use --force). On success prints a JSON summary to stdout.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import expense_common as ec  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, PieChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


def build(stmts, cfg, out_path):
    L = cfg["L"]
    money = ec.money_format(cfg)
    date_format = cfg["date_format"]

    def tx(column):
        return ec.tx_range(cfg, column)

    spend, pay, refund = (ec.q(L[key]) for key in ("type_spend", "type_payment", "type_refund"))
    n = len(stmts)

    wb = Workbook()
    dash = wb.active
    dash.title = L["dashboard"]
    ws_p = wb.create_sheet(L["by_period"])
    ws_c = wb.create_sheet(L["by_category"])
    ws_m = wb.create_sheet(L["merchants"])
    ws_t = wb.create_sheet(L["transactions"])
    ws_h = wb.create_sheet(L["how_to"])

    ec.write_transactions(ws_t, stmts, cfg)

    # By statement: one row per period.
    ec.header_row(ws_p, 1, [L["col_period"], L["col_statement_date"], L["col_net"], L["col_credits"], L["col_txns"],
                            L["col_avg_txn"], L["col_closing"], L["col_min_payment"], L["col_due"], L["col_change"]])
    p_last = n + 1
    for r, s in enumerate(stmts, start=2):
        period = f"{tx('period')},$A{r}"
        ec.put(ws_p, r, 1, s["period"], "@", literal=True)
        ec.put(ws_p, r, 2, s["statement_date"], date_format)
        ec.put(ws_p, r, 3, f"=SUMIFS({tx('net')},{period})", money)
        ec.put(ws_p, r, 4, f"=SUMIFS({tx('amount')},{period},{tx('type')},{pay})+SUMIFS({tx('amount')},{period},{tx('type')},{refund})", money)
        ec.put(ws_p, r, 5, f"=COUNTIFS({period},{tx('type')},{spend})", "#,##0")
        ec.put(ws_p, r, 6, f"=IFERROR(SUMIFS({tx('net')},{period},{tx('type')},{spend})/E{r},0)", money)
        ec.put(ws_p, r, 7, s["closing_balance"], money)
        ec.put(ws_p, r, 8, s["minimum_payment"], money)
        ec.put(ws_p, r, 9, s["due_date"], date_format)
        ec.put(ws_p, r, 10, None if r == 2 else f'=IFERROR(C{r}/C{r - 1}-1,"")', "+0.0%;-0.0%;0.0%")
    ec.band(ws_p, 2, p_last, 1, 10)
    ec.flag_change(ws_p, f"J2:J{p_last}")
    ec.set_widths(ws_p, {"A": 12, "B": 14, "C": 18, "D": 18, "E": 12, "F": 16, "G": 18, "H": 18, "I": 14, "J": 14})
    ws_p.freeze_panes = "A2"

    # By category: summary table, then a category x statement matrix.
    P = ec.sheet_ref(L["by_period"])
    periods = f"COUNTA({P}!$A$2:$A${p_last})"
    cats = cfg["categories"]
    c_last = len(cats) + 1
    c_total = c_last + 1
    ec.header_row(ws_c, 1, [L["col_category"], L["col_net"], L["col_share"], L["col_avg_period"],
                            L["col_budget"], L["col_avg_vs_budget"]])
    for r, cat in enumerate(cats, start=2):
        ec.put(ws_c, r, 1, cat["name"], literal=True)
        ec.put(ws_c, r, 2, f"=SUMIFS({tx('net')},{tx('category')},$A{r})", money)
        ec.put(ws_c, r, 3, f"=IFERROR(B{r}/$B${c_total},0)", ec.PCT)
        ec.put(ws_c, r, 4, f"=IFERROR(B{r}/{periods},0)", money)
        ec.put(ws_c, r, 5, cat.get("budget") or 0, money)
        ec.put(ws_c, r, 6, f"=IFERROR(D{r}/E{r},0)", ec.PCT)
    ec.band(ws_c, 2, c_last, 1, 6)
    ec.total_row(ws_c, c_total,
                 [L["total"], f"=SUM(B2:B{c_last})", f"=SUM(C2:C{c_last})", f"=SUM(D2:D{c_last})",
                  f"=SUM(E2:E{c_last})", f"=IFERROR(D{c_total}/E{c_total},0)"],
                 [None, money, ec.PCT, money, money, ec.PCT])
    ec.flag_over(ws_c, f"F2:F{c_total}", 1)

    m_title = c_total + 2
    m_head = m_title + 1
    m_width = len(cats) + 2
    ec.section_title(ws_c, m_title, L["matrix_title"], m_width)
    ec.header_row(ws_c, m_head, [L["col_period"]] + [c["name"] for c in cats] + [L["total"]])
    for i, s in enumerate(stmts):
        r = m_head + 1 + i
        ec.put(ws_c, r, 1, s["period"], "@", literal=True)
        for j in range(2, len(cats) + 2):
            letter = get_column_letter(j)
            ec.put(ws_c, r, j, f"=SUMIFS({tx('net')},{tx('period')},$A{r},{tx('category')},{letter}${m_head})", money)
        ec.put(ws_c, r, m_width, f"=SUM(B{r}:{get_column_letter(m_width - 1)}{r})", money)
    m_last = m_head + n
    ec.band(ws_c, m_head + 1, m_last, 1, m_width)
    ec.set_widths(ws_c, {get_column_letter(c): 18 for c in range(1, max(m_width, 6) + 1)})
    ec.set_widths(ws_c, {"A": 26})

    # Merchants: groups by total net spend, with one column per statement.
    groups = ec.sorted_groups(stmts)
    ec.header_row(ws_m, 1, [L["col_group"], L["col_txns"], L["col_net"], L["col_share"], L["col_avg_txn"]]
                  + [s["period"] for s in stmts])
    g_last = len(groups) + 1
    for r, group in enumerate(groups, start=2):
        match = f"{tx('group')},$A{r}"
        ec.put(ws_m, r, 1, group, literal=True)
        ec.put(ws_m, r, 2, f"=COUNTIFS({match},{tx('type')},{spend})", "#,##0")
        ec.put(ws_m, r, 3, f"=SUMIFS({tx('net')},{match})", money)
        ec.put(ws_m, r, 4, f"=IFERROR(C{r}/SUM($C$2:$C${g_last}),0)", ec.PCT)
        ec.put(ws_m, r, 5, f"=IFERROR(SUMIFS({tx('net')},{match},{tx('type')},{spend})/B{r},0)", money)
        for j in range(6, 6 + n):
            letter = get_column_letter(j)
            ec.put(ws_m, r, j, f"=SUMIFS({tx('net')},{match},{tx('period')},{letter}$1)", money)
    if groups:
        ec.band(ws_m, 2, g_last, 1, 5 + n)
    ec.set_widths(ws_m, {get_column_letter(c): 16 for c in range(2, 6 + n)})
    ec.set_widths(ws_m, {"A": 30, "B": 12, "D": 10})
    ws_m.freeze_panes = "B2"

    # Dashboard.
    C = ec.sheet_ref(L["by_category"])

    def col(letter):
        return f"{P}!${letter}$2:${letter}${p_last}"

    accounts = sorted({s["account_label"] or s["institution"] for s in stmts} - {""})
    subtitle = " · ".join(accounts + [f"{stmts[0]['period']} → {stmts[-1]['period']}", f"{n} {L['statements_word']}"])
    ec.title_block(dash, cfg.get("overview_title") or L["master_title"], subtitle, 4)
    ec.kpi_row(dash, 5, [
        (L["kpi_net_spend"], f"=SUM({col('C')})", money),
        (L["kpi_avg_period"], f"=IFERROR(SUM({col('C')})/COUNTA({col('A')}),0)", money),
        (L["kpi_txns"], f"=SUM({col('E')})", "#,##0"),
        (L["kpi_credits"], f"=SUM({col('D')})", money),
    ])
    ec.kpi_row(dash, 8, [
        (L["kpi_highest"], f'=INDEX({col("A")},MATCH(MAX({col("C")}),{col("C")},0))&" · "&{ec.text_amount("MAX(" + col("C") + ")", cfg)}', "@"),
        (L["kpi_top_category"], f"=INDEX({C}!$A$2:$A${c_last},MATCH(MAX({C}!$B$2:$B${c_last}),{C}!$B$2:$B${c_last},0))", "@"),
        (L["kpi_latest_closing"], f'=IFERROR(INDEX({col("G")},COUNTA({col("A")})),"")', money),
        (L["kpi_latest_due"], f'=IFERROR(INDEX({col("I")},COUNTA({col("A")})),"")', date_format),
    ])

    by_period = BarChart()
    by_period.type = "col"
    by_period.title = L["chart_period"]
    by_period.add_data(Reference(ws_p, min_col=3, max_col=4, min_row=1, max_row=p_last), titles_from_data=True)
    by_period.set_categories(Reference(ws_p, min_col=1, min_row=2, max_row=p_last))
    by_period.y_axis.numFmt = "#,##0"
    by_period.width, by_period.height = 11.5, 8.5
    dash.add_chart(by_period, "A11")

    pie = PieChart()
    pie.title = L["chart_category"]
    pie.add_data(Reference(ws_c, min_col=2, min_row=1, max_row=c_last), titles_from_data=True)
    pie.set_categories(Reference(ws_c, min_col=1, min_row=2, max_row=c_last))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.width, pie.height = 11.5, 8.5
    dash.add_chart(pie, "C11")

    mix = BarChart()
    mix.type = "col"
    mix.grouping = "stacked"
    mix.overlap = 100
    mix.title = L["chart_mix"]
    mix.add_data(Reference(ws_c, min_col=2, max_col=len(cats) + 1, min_row=m_head, max_row=m_last), titles_from_data=True)
    mix.set_categories(Reference(ws_c, min_col=1, min_row=m_head + 1, max_row=m_last))
    mix.y_axis.numFmt = "#,##0"
    mix.width, mix.height = 23, 9
    dash.add_chart(mix, "A30")

    if groups:
        top = min(10, len(groups))
        merchants = BarChart()
        merchants.type = "bar"
        merchants.title = L["chart_merchants"]
        merchants.add_data(Reference(ws_m, min_col=3, min_row=1, max_row=top + 1), titles_from_data=True)
        merchants.set_categories(Reference(ws_m, min_col=1, min_row=2, max_row=top + 1))
        merchants.x_axis.scaling.orientation = "maxMin"
        merchants.legend = None
        merchants.width, merchants.height = 23, 10
        dash.add_chart(merchants, "A50")

    ec.set_widths(dash, {letter: 32 for letter in "ABCD"})
    dash.sheet_view.showGridLines = False

    # How to update (no merged cells).
    head = ec.put(ws_h, 1, 1, L["how_to"], literal=True)
    head.font = Font(bold=True, size=14, color=ec.NAVY)
    for i, line in enumerate(L["how_to_lines"], start=3):
        cell = ec.put(ws_h, i, 1, line, literal=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ec.set_widths(ws_h, {"A": 120})
    ws_h.sheet_view.showGridLines = False

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    ec.utf8_stdio()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("inputs", nargs="+", help="statement JSON files, folders, or glob patterns")
    parser.add_argument("--config", help="expense-config.json (defaults apply when omitted)")
    parser.add_argument("--out-dir", default=".", help="folder for the overview workbook (default: current folder)")
    parser.add_argument("--allow-unreconciled", action="store_true", help="build even if some totals do not match")
    parser.add_argument("--uncategorized-to-other", action="store_true", help="put unmatched transactions in other_category")
    parser.add_argument("--force", action="store_true", help="overwrite an existing overview workbook")
    args = parser.parse_args()

    try:
        cfg = ec.load_config(args.config)
        paths = ec.expand_inputs(args.inputs)
        if not paths:
            raise ec.InputError("no statement JSON files found")
        stmts, missing = [], []
        for path in paths:
            stmt = ec.load_statement(path)
            missing += [(stmt["period"], t) for t in ec.resolve(stmt, cfg, args.uncategorized_to_other)]
            stmts.append(stmt)
    except ec.InputError as exc:
        ec.fail(ec.EXIT_BAD_INPUT, f"INVALID INPUT: {exc}")

    stmts.sort(key=lambda s: s["period"])
    periods = [s["period"] for s in stmts]
    duplicates = sorted({p for p in periods if periods.count(p) > 1})
    if duplicates:
        ec.fail(ec.EXIT_BAD_INPUT, f"INVALID INPUT: more than one statement for period(s) {', '.join(duplicates)}")
    if missing:
        lines = [f"  {p} #{t['no']} {t['transaction_date']} {t['description']} {t['amount']}" for p, t in missing]
        ec.fail(ec.EXIT_UNCATEGORIZED, f"UNCATEGORIZED ({len(missing)}): set 'category' in the JSON or add config rules\n" + "\n".join(lines))
    issues = {s["period"]: ec.reconcile(s, cfg) for s in stmts}
    issues = {p: v for p, v in issues.items() if v}
    if issues and not args.allow_unreconciled:
        ec.fail(ec.EXIT_UNRECONCILED, "UNRECONCILED:\n" + "\n".join(f"  {p}: {'; '.join(v)}" for p, v in issues.items()))

    out_path = Path(args.out_dir) / f"{cfg['file_prefix']}_{cfg['L']['file_overview']}.xlsx"
    if out_path.exists() and not args.force:
        ec.fail(ec.EXIT_EXISTS, f"EXISTS: {out_path} (re-run with --force to overwrite)")
    build(stmts, cfg, out_path)

    result = {"workbook": str(out_path), "reconciliation": issues or "ok", **ec.summary(stmts, cfg)}
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
