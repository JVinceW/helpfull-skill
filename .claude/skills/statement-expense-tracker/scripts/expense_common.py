"""Shared helpers for the statement-expense-tracker scripts.

The normalized statement JSON (schema v1) is documented in references/mapping.md.
"""
from __future__ import annotations

import collections
import copy
import datetime as dt
import glob
import json
import re
import sys
from pathlib import Path

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SCHEMA_VERSION = 1
TYPES = ("spend", "fee", "refund", "payment")
NET_SIGN = {"spend": 1, "fee": 1, "refund": -1, "payment": 0}
TX_LAST_ROW = 10000  # formulas cover Transactions rows 2..TX_LAST_ROW
PERIOD_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
EXIT_BAD_INPUT, EXIT_UNRECONCILED, EXIT_UNCATEGORIZED, EXIT_EXISTS = 1, 2, 3, 4

# Transactions sheet columns, shared by the monthly and overview workbooks.
TX_COLUMNS = ("period", "no", "txn_date", "post_date", "description", "group",
              "category", "type", "amount", "net", "original", "note")
COL = {name: get_column_letter(i) for i, name in enumerate(TX_COLUMNS, start=1)}

LABELS = {
    "en": {
        "dashboard": "Dashboard", "transactions": "Transactions", "budget": "Budget",
        "merchants": "Merchants", "statement": "Statement", "by_period": "By statement",
        "by_category": "By category", "how_to": "How to update",
        "monthly_title": "EXPENSE DASHBOARD", "master_title": "EXPENSE OVERVIEW - ALL STATEMENTS",
        "file_overview": "Overview",
        "col_period": "Period", "col_no": "#", "col_txn_date": "Transaction date",
        "col_post_date": "Posting date", "col_description": "Description", "col_group": "Merchant group",
        "col_category": "Category", "col_type": "Type", "col_amount": "Amount", "col_net": "Net spend",
        "col_original": "Original amount", "col_note": "Note",
        "type_spend": "Spend", "type_fee": "Fee / interest", "type_refund": "Refund", "type_payment": "Payment",
        "payment_category": "(Payment)",
        "col_budget": "Budget / month", "col_actual": "Actual", "col_remaining": "Remaining", "col_used": "Used",
        "col_txns": "Transactions", "col_share": "Share", "col_avg_txn": "Avg / transaction", "total": "TOTAL",
        "col_statement_date": "Statement date", "col_credits": "Payments & credits",
        "col_closing": "Closing balance", "col_min_payment": "Minimum payment", "col_due": "Due date",
        "col_change": "Change vs previous", "col_avg_period": "Avg / statement", "col_avg_vs_budget": "Avg vs budget",
        "matrix_title": "NET SPEND BY CATEGORY AND STATEMENT",
        "kpi_net_spend": "NET SPEND", "kpi_credits": "PAYMENTS & CREDITS", "kpi_txns": "SPEND TRANSACTIONS",
        "kpi_avg_txn": "AVG / TRANSACTION", "kpi_closing": "CLOSING BALANCE", "kpi_min_payment": "MINIMUM PAYMENT",
        "kpi_due": "PAYMENT DUE", "kpi_top_category": "TOP CATEGORY", "kpi_avg_period": "AVG / STATEMENT",
        "kpi_highest": "HIGHEST STATEMENT", "kpi_latest_closing": "LATEST CLOSING BALANCE",
        "kpi_latest_due": "LATEST DUE DATE",
        "chart_category": "Net spend by category", "chart_budget": "Budget vs actual",
        "chart_merchants": "Top 10 merchants", "chart_period": "Net spend and payments by statement",
        "chart_mix": "Category mix by statement",
        "field": "Field", "value": "Value",
        "f_period": "Period", "f_institution": "Institution", "f_account": "Account", "f_currency": "Currency",
        "f_statement_date": "Statement date", "f_due_date": "Payment due date", "f_opening": "Opening balance",
        "f_debits": "Total debits", "f_credits": "Total credits", "f_closing": "Closing balance",
        "f_min_payment": "Minimum payment", "f_credit_limit": "Credit limit", "f_reconciliation": "Reconciliation",
        "f_generated": "Generated at",
        "reconciled": "OK - transactions match the statement totals",
        "unreconciled": "NOT RECONCILED - {issues}",
        "statements_word": "statements",
        "how_to_lines": [
            "This workbook is generated from data/statement_*.json. Re-run the generator instead of adding rows by hand.",
            "1. Save the new statement as data/statement_YYYY-MM.json (normalized statement, schema v1).",
            "2. Build its workbook: python build_monthly.py data/statement_YYYY-MM.json --config expense-config.json --out-dir .",
            "3. Rebuild this overview: python build_master.py data --config expense-config.json --out-dir . --force",
            "4. To recategorize, edit the rules in expense-config.json or the category field in the JSON, then repeat steps 2-3.",
        ],
    },
    "vi": {
        "dashboard": "Dashboard", "transactions": "Giao dịch", "budget": "Ngân sách",
        "merchants": "Merchant", "statement": "Sao kê", "by_period": "Theo kỳ",
        "by_category": "Theo danh mục", "how_to": "Hướng dẫn",
        "monthly_title": "BÁO CÁO CHI TIÊU", "master_title": "TỔNG HỢP CHI TIÊU - TẤT CẢ CÁC KỲ",
        "file_overview": "Tong_hop",
        "col_period": "Kỳ sao kê", "col_no": "#", "col_txn_date": "Ngày giao dịch",
        "col_post_date": "Ngày hạch toán", "col_description": "Mô tả", "col_group": "Nhóm merchant",
        "col_category": "Danh mục", "col_type": "Loại", "col_amount": "Số tiền", "col_net": "Chi tiêu ròng",
        "col_original": "Số tiền gốc", "col_note": "Ghi chú",
        "type_spend": "Chi tiêu", "type_fee": "Phí / lãi", "type_refund": "Hoàn tiền", "type_payment": "Thanh toán",
        "payment_category": "(Thanh toán)",
        "col_budget": "Ngân sách / tháng", "col_actual": "Thực chi", "col_remaining": "Còn lại", "col_used": "Đã dùng",
        "col_txns": "Số GD", "col_share": "Tỷ trọng", "col_avg_txn": "TB / giao dịch", "total": "TỔNG",
        "col_statement_date": "Ngày sao kê", "col_credits": "Thanh toán & hoàn tiền",
        "col_closing": "Dư nợ cuối kỳ", "col_min_payment": "Thanh toán tối thiểu", "col_due": "Hạn thanh toán",
        "col_change": "So với kỳ trước", "col_avg_period": "TB / kỳ", "col_avg_vs_budget": "TB so với ngân sách",
        "matrix_title": "CHI TIÊU RÒNG THEO DANH MỤC VÀ KỲ SAO KÊ",
        "kpi_net_spend": "CHI TIÊU RÒNG", "kpi_credits": "THANH TOÁN & HOÀN TIỀN", "kpi_txns": "SỐ GIAO DỊCH CHI",
        "kpi_avg_txn": "TB / GIAO DỊCH", "kpi_closing": "DƯ NỢ CUỐI KỲ", "kpi_min_payment": "THANH TOÁN TỐI THIỂU",
        "kpi_due": "HẠN THANH TOÁN", "kpi_top_category": "DANH MỤC LỚN NHẤT", "kpi_avg_period": "TB / KỲ",
        "kpi_highest": "KỲ CHI CAO NHẤT", "kpi_latest_closing": "DƯ NỢ KỲ GẦN NHẤT",
        "kpi_latest_due": "HẠN THANH TOÁN GẦN NHẤT",
        "chart_category": "Chi tiêu theo danh mục", "chart_budget": "Ngân sách và thực chi",
        "chart_merchants": "Top 10 merchant", "chart_period": "Chi tiêu và thanh toán theo kỳ",
        "chart_mix": "Cơ cấu danh mục theo kỳ",
        "field": "Mục", "value": "Giá trị",
        "f_period": "Kỳ sao kê", "f_institution": "Tổ chức phát hành", "f_account": "Tài khoản", "f_currency": "Tiền tệ",
        "f_statement_date": "Ngày sao kê", "f_due_date": "Hạn thanh toán", "f_opening": "Dư nợ đầu kỳ",
        "f_debits": "Phát sinh nợ", "f_credits": "Phát sinh có", "f_closing": "Dư nợ cuối kỳ",
        "f_min_payment": "Thanh toán tối thiểu", "f_credit_limit": "Hạn mức tín dụng", "f_reconciliation": "Đối soát",
        "f_generated": "Thời điểm tạo",
        "reconciled": "OK - giao dịch khớp với tổng trên sao kê",
        "unreconciled": "CHƯA KHỚP - {issues}",
        "statements_word": "kỳ",
        "how_to_lines": [
            "Workbook này được tạo tự động từ data/statement_*.json. Hãy chạy lại script thay vì thêm dòng thủ công.",
            "1. Lưu sao kê mới thành data/statement_YYYY-MM.json (định dạng chuẩn hóa, schema v1).",
            "2. Tạo workbook của kỳ: python build_monthly.py data/statement_YYYY-MM.json --config expense-config.json --out-dir .",
            "3. Tạo lại file tổng hợp: python build_master.py data --config expense-config.json --out-dir . --force",
            "4. Muốn đổi danh mục: sửa rules trong expense-config.json hoặc trường category trong JSON, rồi làm lại bước 2-3.",
        ],
    },
}

DEFAULT_CONFIG = {
    "language": "en",
    "currency": {"code": "USD", "symbol": "$", "symbol_position": "prefix", "decimals": 2},
    "date_format": "yyyy-mm-dd",
    "file_prefix": "Expenses",
    "other_category": "Other",
    "categories": [],
    "rules": [],
    "merchant_groups": [],
    "strip_prefixes": [],
    "labels": {},
}

NAVY, BLUE, LIGHT, BAND, GRID = "1F4E78", "4A90C8", "DDEBF7", "F3F8FC", "BFD7EA"
PCT = "0.0%"
_thin = Side(style="thin", color=GRID)
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


class InputError(Exception):
    """Invalid config or statement input."""


def utf8_stdio():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass


def fail(code, message):
    print(message, file=sys.stderr)
    sys.exit(code)


# ---------------------------------------------------------------- config

def load_config(path):
    cfg = copy.deepcopy(DEFAULT_CONFIG)
    if path:
        try:
            user = json.loads(Path(path).read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise InputError(f"config {path}: {exc}") from exc
        for key, value in user.items():
            if key == "currency":
                cfg["currency"] = {"code": "", "symbol": "", "symbol_position": "suffix", "decimals": 2, **value}
            else:
                cfg[key] = value
    if cfg["language"] not in LABELS:
        raise InputError(f"config: language must be one of {sorted(LABELS)}")
    names = [c.get("name") for c in cfg["categories"]]
    if not all(names) or len(set(names)) != len(names):
        raise InputError("config: categories need unique, non-empty names")
    if cfg["other_category"] not in names:
        cfg["categories"].append({"name": cfg["other_category"], "budget": 0})
        names.append(cfg["other_category"])
    for rule in cfg["rules"]:
        if not rule.get("keyword") or rule.get("category") not in names:
            raise InputError(f"config: rule {rule!r} needs a keyword and a category listed in categories")
    for group in cfg["merchant_groups"]:
        if not group.get("keyword") or not group.get("name"):
            raise InputError(f"config: merchant group {group!r} needs keyword and name")
    labels = copy.deepcopy(LABELS[cfg["language"]])
    labels.update(cfg.get("labels") or {})
    cfg["L"] = labels
    return cfg


# ---------------------------------------------------------------- statements

def _date(value, field, required=False):
    if value in (None, ""):
        if required:
            raise InputError(f"{field} is required (YYYY-MM-DD)")
        return None
    try:
        return dt.date.fromisoformat(str(value))
    except ValueError as exc:
        raise InputError(f"{field}: expected YYYY-MM-DD, got {value!r}") from exc


def _number(value, field, required=False):
    if value is None:
        if required:
            raise InputError(f"{field} is required")
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise InputError(f"{field}: expected a number, got {value!r}")
    return value


def load_statement(path):
    path = Path(path)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise InputError(f"{path}: {exc}") from exc
    where = path.name
    if data.get("schema_version") != SCHEMA_VERSION:
        raise InputError(f"{where}: schema_version must be {SCHEMA_VERSION}")
    period = str(data.get("period", ""))
    if not PERIOD_RE.match(period):
        raise InputError(f"{where}: period must be YYYY-MM, got {period!r}")
    account = str(data.get("account_label") or "")
    if re.search(r"\d{7,}", re.sub(r"[\s-]", "", account)):
        raise InputError(f"{where}: account_label looks like a full account number; mask it, e.g. 'Visa ****1234'")
    s = data.get("statement") or {}
    stmt = {
        "period": period,
        "institution": str(data.get("institution") or ""),
        "account_label": account,
        "currency": str(data.get("currency") or ""),
        "statement_date": _date(s.get("statement_date"), f"{where}: statement.statement_date", required=True),
        "due_date": _date(s.get("due_date"), f"{where}: statement.due_date"),
    }
    for key in ("opening_balance", "total_debits", "total_credits", "closing_balance", "minimum_payment", "credit_limit"):
        stmt[key] = _number(s.get(key), f"{where}: statement.{key}")
    txs = []
    for index, raw in enumerate(data.get("transactions") or []):
        field = f"{where}: transactions[{index}]"
        kind = raw.get("type")
        if kind not in TYPES:
            raise InputError(f"{field}.type must be one of {', '.join(TYPES)}, got {kind!r}")
        amount = _number(raw.get("amount"), f"{field}.amount", required=True)
        if amount < 0:
            raise InputError(f"{field}.amount must be positive; use type refund or payment for credits")
        description = str(raw.get("description") or "").strip()
        if not description:
            raise InputError(f"{field}.description is required")
        original = ""
        if raw.get("original_amount") is not None:
            original = f"{raw['original_amount']} {raw.get('original_currency') or ''}".strip()
        txs.append({
            "no": index + 1,
            "transaction_date": _date(raw.get("transaction_date"), f"{field}.transaction_date", required=True),
            "posting_date": _date(raw.get("posting_date"), f"{field}.posting_date"),
            "description": description,
            "type": kind,
            "amount": amount,
            "net": amount * NET_SIGN[kind],
            "original": original,
            "category": raw.get("category") or None,
            "note": str(raw.get("note") or ""),
        })
    if not txs:
        raise InputError(f"{where}: transactions is empty")
    stmt["transactions"] = txs
    return stmt


def expand_inputs(items):
    paths = []
    for item in items:
        p = Path(item)
        if p.is_dir():
            paths.extend(sorted(p.glob("*.json")))
        elif any(ch in item for ch in "*?["):
            paths.extend(Path(x) for x in sorted(glob.glob(item)))
        else:
            paths.append(p)
    return paths


def _match(description, entries, key):
    upper = description.upper()
    for entry in entries:
        if str(entry["keyword"]).upper() in upper:
            return entry[key]
    return None


def merchant_group(description, cfg):
    group = _match(description, cfg["merchant_groups"], "name")
    if group:
        return group
    text = description
    for prefix in cfg["strip_prefixes"]:
        if text.upper().startswith(prefix.upper()):
            text = text[len(prefix):]
            break
    return text.strip() or description


def resolve(stmt, cfg, uncategorized_to_other=False):
    """Fill category and merchant group; return transactions that are still uncategorized."""
    names = {c["name"] for c in cfg["categories"]}
    missing = []
    for tx in stmt["transactions"]:
        if tx["type"] == "payment":
            tx["group"] = tx["category"] = cfg["L"]["payment_category"]
            continue
        tx["group"] = merchant_group(tx["description"], cfg)
        if tx["category"]:
            if tx["category"] not in names:
                raise InputError(f"{stmt['period']} #{tx['no']}: category {tx['category']!r} is not in config categories")
            continue
        tx["category"] = _match(tx["description"], cfg["rules"], "category")
        if tx["category"] is None:
            if uncategorized_to_other:
                tx["category"] = cfg["other_category"]
            else:
                missing.append(tx)
    return missing


def decimals(cfg):
    return int(cfg["currency"].get("decimals", 2))


def _plain(value, places):
    return f"{value:,.{places}f}"


def reconcile(stmt, cfg):
    """Return a list of human-readable reconciliation problems (empty when totals match)."""
    places = decimals(cfg)
    tolerance = 0.5 * 10 ** -places + 1e-9
    txs = stmt["transactions"]
    debits = round(sum(t["amount"] for t in txs if t["type"] in ("spend", "fee")), places)
    credits = round(sum(t["amount"] for t in txs if t["type"] in ("refund", "payment")), places)
    issues = []

    def check(name, expected, actual):
        if expected is not None and abs(expected - actual) > tolerance:
            issues.append(f"{name}: statement {_plain(expected, places)} vs computed {_plain(actual, places)}")

    if stmt["total_debits"] is None and stmt["total_credits"] is None:
        issues.append("statement.total_debits and statement.total_credits are missing, nothing can be reconciled")
    check("total_debits", stmt["total_debits"], debits)
    check("total_credits", stmt["total_credits"], credits)
    parts = (stmt["opening_balance"], stmt["total_debits"], stmt["total_credits"], stmt["closing_balance"])
    if None not in parts:
        check("closing_balance", parts[3], round(parts[0] + parts[1] - parts[2], places))
    return issues


def summary(stmts, cfg):
    places = decimals(cfg)

    def rnd(value):
        return None if value is None else round(value, places)

    per, cats, groups = {}, collections.Counter(), collections.Counter()
    for s in stmts:
        txs = s["transactions"]
        per[s["period"]] = {
            "net_spend": rnd(sum(t["net"] for t in txs)),
            "payments_and_credits": rnd(sum(t["amount"] for t in txs if t["type"] in ("payment", "refund"))),
            "spend_transactions": sum(1 for t in txs if t["type"] == "spend"),
            "closing_balance": s["closing_balance"],
            "minimum_payment": s["minimum_payment"],
            "due_date": s["due_date"].isoformat() if s["due_date"] else None,
        }
        for t in txs:
            if t["type"] != "payment":
                cats[t["category"]] += t["net"]
                groups[t["group"]] += t["net"]
    count = len(stmts)
    categories = {}
    for c in cfg["categories"]:
        total = cats.get(c["name"], 0)
        budget = c.get("budget") or 0
        categories[c["name"]] = {"net_spend": rnd(total), "avg_per_statement": rnd(total / count),
                                 "budget": budget, "over_budget": bool(budget) and total / count > budget}
    top = sorted(groups.items(), key=lambda kv: (-kv[1], kv[0]))[:10]
    return {"statements": per, "total_net_spend": rnd(sum(cats.values())), "categories": categories,
            "top_merchants": [[g, rnd(v)] for g, v in top]}


def sorted_groups(stmts):
    totals = collections.Counter()
    for s in stmts:
        for t in s["transactions"]:
            if t["type"] != "payment":
                totals[t["group"]] += t["net"]
    return sorted(totals, key=lambda g: (-totals[g], g))


# ---------------------------------------------------------------- workbook helpers

def money_format(cfg):
    places = decimals(cfg)
    number = "#,##0" + ("." + "0" * places if places else "")
    symbol = cfg["currency"].get("symbol") or ""
    if not symbol:
        return number
    if cfg["currency"].get("symbol_position") == "prefix":
        return f'"{symbol}"{number}'
    return f'{number} "{symbol}"'


def text_amount(expr, cfg):
    places = decimals(cfg)
    body = f'TEXT({expr},"#,##0{"." + "0" * places if places else ""}")'
    symbol = (cfg["currency"].get("symbol") or "").replace('"', '""')
    if not symbol:
        return body
    if cfg["currency"].get("symbol_position") == "prefix":
        return f'"{symbol}"&{body}'
    return f'{body}&" {symbol}"'


def q(text):
    """Quote a string literal for use inside a formula."""
    return '"' + str(text).replace('"', '""') + '"'


def sheet_ref(name):
    return "'" + name.replace("'", "''") + "'"


def tx_range(cfg, column):
    letter = COL[column]
    return f"{sheet_ref(cfg['L']['transactions'])}!${letter}$2:${letter}${TX_LAST_ROW}"


def put(ws, row, col, value, fmt=None, literal=False):
    cell = ws.cell(row=row, column=col, value=value)
    if literal and isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    if fmt:
        cell.number_format = fmt
    return cell


def header_row(ws, row, values, col=1):
    for offset, value in enumerate(values):
        cell = put(ws, row, col + offset, value, "@", literal=True)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 30


def band(ws, first_row, last_row, first_col, last_col):
    for r in range(first_row, last_row + 1):
        for c in range(first_col, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            if (r - first_row) % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=BAND)


def total_row(ws, row, values, formats):
    for col, (value, fmt) in enumerate(zip(values, formats), start=1):
        cell = put(ws, row, col, value, fmt)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.border = BOX


def set_widths(ws, widths):
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def title_block(ws, title, subtitle, last_col):
    # Merges end in a column that holds values below, so Google Sheets' xlsx import keeps them.
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=last_col)
    cell = put(ws, 1, 1, title, literal=True)
    cell.font = Font(bold=True, size=16, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if subtitle:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
        cell = put(ws, 3, 1, subtitle, literal=True)
        cell.font = Font(italic=True, color="555555")
        cell.alignment = Alignment(horizontal="center")


def section_title(ws, row, text, last_col):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = put(ws, row, 1, text, literal=True)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(horizontal="center")


def kpi_row(ws, row, cards):
    """One KPI card per column: label at `row`, value at `row + 1` (no merged cells)."""
    for col, (label, formula, fmt) in enumerate(cards, start=1):
        head = put(ws, row, col, label, literal=True)
        head.font = Font(bold=True, size=10, color="FFFFFF")
        head.fill = PatternFill("solid", fgColor=BLUE)
        head.alignment = Alignment(horizontal="center", vertical="center")
        value = put(ws, row + 1, col, formula, fmt)
        value.font = Font(bold=True, size=14, color=NAVY)
        value.fill = PatternFill("solid", fgColor=LIGHT)
        value.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row + 1].height = 30


def flag_over(ws, ref, threshold):
    ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=[str(threshold)],
                                                  font=Font(color="C00000", bold=True),
                                                  fill=PatternFill("solid", fgColor="F8D7DA")))


def flag_negative(ws, ref):
    ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000")))


def flag_change(ws, ref):
    ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="C00000", bold=True)))
    ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="2E7D32", bold=True)))


def write_transactions(ws, stmts, cfg):
    L = cfg["L"]
    money = money_format(cfg)
    date_format = cfg["date_format"]
    header_row(ws, 1, [L["col_" + key] for key in ("period", "no", "txn_date", "post_date", "description", "group",
                                                    "category", "type", "amount", "net", "original", "note")])
    row = 1
    for s in stmts:
        for t in s["transactions"]:
            row += 1
            put(ws, row, 1, s["period"], "@", literal=True)
            put(ws, row, 2, t["no"])
            put(ws, row, 3, t["transaction_date"], date_format)
            put(ws, row, 4, t["posting_date"], date_format)
            put(ws, row, 5, t["description"], literal=True)
            put(ws, row, 6, t["group"], literal=True)
            put(ws, row, 7, t["category"], literal=True)
            put(ws, row, 8, L["type_" + t["type"]], literal=True)
            put(ws, row, 9, t["amount"], money)
            put(ws, row, 10, t["net"], money)
            put(ws, row, 11, t["original"] or None, literal=True)
            put(ws, row, 12, t["note"] or None, literal=True)
    band(ws, 2, row, 1, len(TX_COLUMNS))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{COL['note']}{row}"
    set_widths(ws, {"A": 11, "B": 6, "C": 14, "D": 14, "E": 34, "F": 24, "G": 22, "H": 14,
                    "I": 16, "J": 16, "K": 18, "L": 30})
    return row
